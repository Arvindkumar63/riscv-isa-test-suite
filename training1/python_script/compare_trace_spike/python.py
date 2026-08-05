import sys
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.page import PageMargins

# ------------------------------------------------------------
# ARGUMENT CHECK
# ------------------------------------------------------------
if len(sys.argv) != 4:
    print(f"Usage: python3 {sys.argv[0]} <trace_rtl.log> <spike_tb.log> <output_excel>")
    sys.exit(1)

trace_file = sys.argv[1]
spike_file = sys.argv[2]
excel_file = sys.argv[3]

# ------------------------------------------------------------
# REGISTER MAP
# ------------------------------------------------------------
reg_map = {
"ra":"x1","sp":"x2","gp":"x3","tp":"x4",
"t0":"x5","t1":"x6","t2":"x7",
"s0":"x8","fp":"x8","s1":"x9",
"a0":"x10","a1":"x11","a2":"x12","a3":"x13",
"a4":"x14","a5":"x15","a6":"x16","a7":"x17",
"s2":"x18","s3":"x19","s4":"x20","s5":"x21",
"s6":"x22","s7":"x23","s8":"x24","s9":"x25",
"s10":"x26","s11":"x27",
"t3":"x28","t4":"x29","t5":"x30","t6":"x31"
}

# ------------------------------------------------------------
# PARSE trace_rtl.log
# ------------------------------------------------------------
trace_data = []

with open(trace_file) as f:
    for i,line in enumerate(f):

        parts = line.strip().split()

        if len(parts) < 7:
            continue

        time = parts[0]
        cycle = parts[1]
        mode = parts[2]
        pc = "0x"+parts[3].lower()
        opcode = parts[5]

        mnemonic = parts[6]

        operands = ""
        if len(parts) > 7:
            operands = parts[7]

        instr = mnemonic + " " + operands

        rd = ""
        value = ""

        if ":" in line:
            try:
                reg_part = line.split(":")[0].split()[-1]
                val_part = line.split(":")[1].strip()

                rd = reg_map.get(reg_part,reg_part)
                value = "0x"+val_part.lower()

            except:
                pass

        trace_data.append({
            "sl.no":i+1,
            "time":time,
            "cycle":cycle,
            "mode":mode,
            "pc":pc,
            "opcode":opcode,
            "instruction":instr.strip(),
            "pneumonics":mnemonic,
            "rd":rd,
            "value":value
        })

trace_df = pd.DataFrame(trace_data)

# ------------------------------------------------------------
# PARSE spike_tb.log
# ------------------------------------------------------------
spike_pattern = re.compile(
r"core\s+\d+:\s+\d+\s+(0x[0-9a-fA-F]+)\s+\((0x[0-9a-fA-F]+)\)\s*(x\s*\d+)?\s*(0x[0-9a-fA-F]+)?"
)

spike_data = []

with open(spike_file) as f:

    for line in f:

        m = spike_pattern.search(line)

        if not m:
            continue

        pc = m.group(1).lower()
        opcode = m.group(2).replace("0x","")

        rd = ""
        val = ""

        if m.group(3):
            rd = m.group(3).replace(" ","")

        if m.group(4):
            val = m.group(4).lower()

        spike_data.append({
            "pc":pc,
            "opcode":opcode,
            "rd":rd,
            "value":val
        })

spike_df = pd.DataFrame(spike_data)

# ------------------------------------------------------------
# COMPARE
# ------------------------------------------------------------
compare_results = []

min_len = min(len(trace_df), len(spike_df))

for i in range(min_len):   # ? no spaces before this

    t = trace_df.iloc[i]
    s = spike_df.iloc[i]

    print(t)
    print(s)

    pc_match = (t["pc"].lower() == s["pc"].lower())
    rd_match = (t["rd"] == s["rd"])
    val_match = (t["value"].lower() == s["value"].lower())

    status = "PASS" if (pc_match and rd_match and val_match) else "FAIL"
    
    

    compare_results.append({
        "Sl.No":t["sl.no"],
        "Time":t["time"],
        "Cycle":t["cycle"],
        "Mode":t["mode"],
        "PC Value":t["pc"],
        "Opcode":t["opcode"],
        "Instruction":t["instruction"],
        "Pneumonics":t["pneumonics"],
        "RD":t["rd"],
        "Value":t["value"],
        "trac pc ": t["pc"],
        "spike pc ": s["pc"],
        "Compare":status
    })

result_df = pd.DataFrame(compare_results)

# ------------------------------------------------------------
# WRITE EXCEL
# ------------------------------------------------------------
result_df.to_excel(excel_file,index=False)

wb = load_workbook(excel_file)
ws = wb.active

# HEADER FORMAT
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

# AUTO WIDTH
for col in ws.columns:

    max_len = 0
    col_letter = col[0].column_letter

    for cell in col:
        try:
            max_len = max(max_len,len(str(cell.value)))
        except:
            pass

    ws.column_dimensions[col_letter].width = max_len + 3

# PAGE SETTINGS
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.page_margins = PageMargins(left=0.3,right=0.3,top=0.5,bottom=0.5)

wb.save(excel_file)

print(f"[OK] Comparison Excel generated -> {excel_file}")