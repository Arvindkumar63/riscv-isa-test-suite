import os
import sys
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.page import PageMargins

# ------------------------------------------------------------
# ARGUMENT CHECK
# ------------------------------------------------------------
if len(sys.argv) != 3:
    print(f"Usage: python3 {sys.argv[0]} <input_log> <output_dump>")
    sys.exit(1)

inp_file  = sys.argv[1]
out_file  = sys.argv[2]
excel_file = out_file.replace(".dump", ".xlsx")

# ------------------------------------------------------------
# READ INPUT LOG
# ------------------------------------------------------------
with open(inp_file, "r") as f:
    lines = f.readlines()

# ------------------------------------------------------------
# REGEX PATTERNS
# ------------------------------------------------------------
pattern_instr = re.compile(
    r"core\s+\d+:\s+(0x[0-9a-fA-F]+)\s+\((0x[0-9a-fA-F]+)\)\s+(\S+)\s+(.*)"
)

pattern_reg = re.compile(
    r"core\s+\d+:\s+3\s+(0x[0-9a-fA-F]+).*?\(\w+\)\s+x\s*(\d+)\s*(0x[0-9a-fA-F]+)"
)

spike_lines = []
excel_data = []
current_instr = None

# ------------------------------------------------------------
# PARSE LOG
# ------------------------------------------------------------
for line in lines:
    line = line.strip()

    m_instr = pattern_instr.search(line)
    if m_instr:
        addr     = m_instr.group(1).replace("0x", "")
        opcode   = m_instr.group(2).replace("0x", "")
        mnemonic = m_instr.group(3)
        operands = m_instr.group(4).strip()

        current_instr = {
            'address': addr,
            'opcode': opcode,
            'mnemonic': mnemonic,
            'operands': operands
        }
        continue

    m_reg = pattern_reg.search(line)
    if m_reg and current_instr:

        reg_num = m_reg.group(2)
        value   = m_reg.group(3).replace("0x", "").zfill(16)
        reg_name = f"x{reg_num}"

        spike_line = (
            f"{current_instr['address']} "
            f"{current_instr['opcode']} "
            f"{current_instr['mnemonic']:<10} "
            f"{current_instr['operands']:<20} "
            f"{reg_name:<4} "
            f"0x{value}"
        )

        spike_lines.append(spike_line)

        excel_data.append({
            "Address": f"0x{current_instr['address']}",
            "Opcode": current_instr["opcode"],
            "Instruction": f"{current_instr['mnemonic']} {current_instr['operands']}",
            "Register": reg_name,
            "Value": f"0x{value}"
        })

        current_instr = None

# ------------------------------------------------------------
# WRITE spike.dump
# ------------------------------------------------------------
with open(out_file, "w") as f:
    for s in spike_lines:
        f.write(s + "\n")

print(f"[OK] spike dump generated: {out_file}")

# ------------------------------------------------------------
# CREATE EXCEL
# ------------------------------------------------------------
if excel_data:
    df = pd.DataFrame(excel_data)
    df.index = df.index + 1
    df.to_excel(excel_file, index=True, sheet_name="Instructions")

    wb = load_workbook(excel_file)
    ws = wb.active

    # HEADER FORMAT
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # AUTO-FIT COLUMN WIDTHS
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter

        for cell in col:
            try:
                max_length = max(max_length, len(str(cell.value)))
            except:
                pass

        ws.column_dimensions[col_letter].width = max_length + 3

    # PAGE SETTINGS FOR PDF EXPORT
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.5, bottom=0.5)

    wb.save(excel_file)
    print(f"[OK] Excel file generated: {excel_file}")

else:
    print("[INFO] No Excel data found.")
