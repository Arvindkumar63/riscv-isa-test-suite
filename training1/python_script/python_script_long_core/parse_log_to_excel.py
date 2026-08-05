import os
import sys
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

# ------------------------------------------------------------
# ARGUMENT CHECK
# ------------------------------------------------------------
if len(sys.argv) != 3:
    print(f"Usage: python3 {sys.argv[0]} <input_log> <output_dump>")
    sys.exit(1)

inp_file  = sys.argv[1]
out_file  = sys.argv[2]
excel_file = out_file.replace(".dump", ".xlsx")

# ------------------------------------------------------------
# READ INPUT LOG
# ------------------------------------------------------------
with open(inp_file, "r") as f:
    lines = f.readlines()

# ------------------------------------------------------------
# REGEX PATTERNS FOR YOUR LOG FORMAT
# ------------------------------------------------------------

# Example:
# core   0: 0x000000008000010c (0x000040a9) c.li    ra, 10
pattern_instr = re.compile(
    r"core\s+\d+:\s+(0x[0-9a-fA-F]+)\s+\((0x[0-9a-fA-F]+)\)\s+(\S+)\s+(.*)"
)

# Example:
# core   0: 3 0x000000008000010c (0x40a9) x 1 0x000000000000000a
pattern_reg = re.compile(
    r"core\s+\d+:\s+3\s+(0x[0-9a-fA-F]+).*?\(\w+\)\s+x\s*(\d+)\s*(0x[0-9a-fA-F]+)"
)

# ------------------------------------------------------------
# STORAGE
# ------------------------------------------------------------
spike_lines = []
excel_data = []
current_instr = None

# ------------------------------------------------------------
# PARSE LOG
# ------------------------------------------------------------
for line in lines:
    line = line.strip()

    # ----------------------------
    # INSTRUCTION LINE
    # ----------------------------
    m_instr = pattern_instr.search(line)
    if m_instr:
        addr     = m_instr.group(1).replace("0x", "")
        opcode   = m_instr.group(2).replace("0x", "")
        mnemonic = m_instr.group(3)
        operands = m_instr.group(4).strip()

        current_instr = {
            'address': addr,
            'opcode': opcode,
            'mnemonic': mnemonic,
            'operands': operands
        }
        continue

    # ----------------------------
    # REGISTER UPDATE LINE
    # ----------------------------
    m_reg = pattern_reg.search(line)
    if m_reg and current_instr:

        pc_addr = m_reg.group(1)
        reg_num = m_reg.group(2)
        value   = m_reg.group(3).replace("0x", "").zfill(16)

        reg_name = f"x{reg_num}"

        # BUILD spike.dump line
        spike_line = (
            f"{current_instr['address']} "
            f"{current_instr['opcode']} "
            f"{current_instr['mnemonic']:<10} "
            f"{current_instr['operands']:<20} "
            f"{reg_name:<4} "
            f"0x{value}"
        )

        spike_lines.append(spike_line)

        # SAVE FOR EXCEL
        excel_data.append({
            "Address": f"0x{current_instr['address']}",
            "Opcode": current_instr["opcode"],
            "Instruction": f"{current_instr['mnemonic']} {current_instr['operands']}",
            "Register": reg_name,
            "Value": f"0x{value}"
        })

        current_instr = None

# ------------------------------------------------------------
# WRITE spike.dump
# ------------------------------------------------------------
with open(out_file, "w") as f:
    for s in spike_lines:
        f.write(s + "\n")

print(f"[OK] spike dump generated: {out_file}")

# ------------------------------------------------------------
# CREATE EXCEL
# ------------------------------------------------------------
if excel_data:
    df = pd.DataFrame(excel_data)
    df.to_excel(excel_file, index=False, sheet_name="Instructions")

    wb = load_workbook(excel_file)
    ws = wb.active

    # Header formatting
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 20

    wb.save(excel_file)
    print("[OK] Excel file generated: {excel_file}")

else:
    print("[INFO] No Excel data found.")
